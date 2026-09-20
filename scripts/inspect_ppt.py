import pptx

prs = pptx.Presentation('4021104954784895500.pptx')

with open('slide_detailed_analysis.txt', 'w', encoding='utf-8') as out:
    for i, slide in enumerate(prs.slides):
        out.write(f'=== SLIDE {i+1} ===\n')
        for s in slide.shapes:
            if s.has_text_frame:
                tf = s.text_frame
                p_texts = [p.text for p in tf.paragraphs if p.text.strip()]
                font_info = [(p.font.name, p.font.size, p.font.color.rgb if p.font.color and p.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else None) for p in tf.paragraphs if p.text.strip()]
                out.write(f'  TextShape "{s.name}" pos=({s.left},{s.top},{s.width},{s.height}): {repr(p_texts)} info={font_info}\n')
            elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                out.write(f'  Group "{s.name}" pos=({s.left},{s.top},{s.width},{s.height}):\n')
                for c in s.shapes:
                    if c.has_text_frame:
                        p_texts = [p.text for p in c.text_frame.paragraphs if p.text.strip()]
                        font_info = [(p.font.name, p.font.size, p.font.color.rgb if p.font.color and p.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else None) for p in c.text_frame.paragraphs if p.text.strip()]
                        out.write(f'    Child "{c.name}" pos=({c.left},{c.top},{c.width},{c.height}): {repr(p_texts)} info={font_info}\n')
                    else:
                        out.write(f'    ChildShape "{c.name}" type={c.shape_type} pos=({c.left},{c.top},{c.width},{c.height})\n')
            else:
                out.write(f'  Shape "{s.name}" type={s.shape_type} pos=({s.left},{s.top},{s.width},{s.height})\n')

print('Wrote slide_detailed_analysis.txt')
