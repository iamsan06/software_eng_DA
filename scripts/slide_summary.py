import pptx

prs = pptx.Presentation('4021104954784895500.pptx')

print(f"Total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    tbs = []
    groups = []
    freeforms = []
    for s in slide.shapes:
        if s.name.startswith("TextBox") or s.has_text_frame:
            tbs.append(s.name)
        elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            groups.append(s.name)
        else:
            freeforms.append(s.name)
    
    # get text preview
    texts = []
    def get_t(s):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for c in s.shapes:
                get_t(c)
    for s in slide.shapes:
        get_t(s)
    
    first = texts[0] if texts else "EMPTY"
    second = texts[1] if len(texts) > 1 else ""
    print(f"Slide {i+1:02d}: {len(slide.shapes)} shapes | {first[:30]} | {second[:30]}")
