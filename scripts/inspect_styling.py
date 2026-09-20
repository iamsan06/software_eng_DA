import pptx

prs = pptx.Presentation('4021104954784895500.pptx')
for idx in [0, 1, 2, 5, 7, 13, 19, 21, 27, 29, 33, 38, 39]:
    slide = prs.slides[idx]
    print(f'=== Slide {idx+1} ===')
    def check_shape(s):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    r_fonts = [(r.font.name, r.font.size, r.font.bold, r.font.color.rgb if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else None) for r in p.runs]
                    print(f'  {s.name}: "{p.text[:40]}" -> runs: {r_fonts}')
        elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for c in s.shapes:
                check_shape(c)
    for s in slide.shapes:
        check_shape(s)
