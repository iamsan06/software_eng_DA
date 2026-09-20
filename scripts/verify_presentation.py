import pptx

prs = pptx.Presentation('final_scientific_proposal.pptx')

print(f"Total slides: {len(prs.slides)}")
assert len(prs.slides) == 40, "Slide count changed!"

found_sections = []
found_figures = []

for idx, slide in enumerate(prs.slides):
    texts = []
    pics = []
    def extract_info(s):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for c in s.shapes:
                extract_info(c)
        elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pics.append((s.name, s.left, s.top, s.width, s.height))
            
    for s in slide.shapes:
        extract_info(s)
        
    full_text = " ".join(texts)
    if "Abstract" in full_text: found_sections.append("Abstract")
    if "Motivation" in full_text: found_sections.append("Motivation")
    if "Datasets" in full_text or "Dataset" in full_text: found_sections.append("Datasets")
    if "Preprocessing" in full_text: found_sections.append("Preprocessing")
    if "Experimental" in full_text or "Benchmark" in full_text: found_sections.append("Experiments")
    if "Conclusion" in full_text: found_sections.append("Conclusion")
    if "Future" in full_text: found_sections.append("Future works")
    if "References" in full_text: found_sections.append("References")
    
    if pics:
        found_figures.append((idx+1, len(pics)))
        print(f"Slide {idx+1:02d} has {len(pics)} PICTURE(s): {pics}")
    else:
        preview = " | ".join(texts[:2])
        print(f"Slide {idx+1:02d}: {preview[:85]}")

print("\n--- Summary of Verified Sections ---")
for sec in ["Abstract", "Motivation", "Datasets", "Preprocessing", "Experiments", "Conclusion", "Future works", "References"]:
    print(f"Section '{sec}': {'FOUND' if sec in found_sections else 'MISSING'}")

print(f"\nTotal slides with embedded pictures: {len(found_figures)} slides ({sum([count for _, count in found_figures])} figures total)")
