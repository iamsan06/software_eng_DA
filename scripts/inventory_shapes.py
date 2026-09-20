import pptx

prs = pptx.Presentation('4021104954784895500.pptx')

with open('slide_shape_inventory.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"=== SLIDE {idx+1} ===\n")
        for s in slide.shapes:
            if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                f.write(f"  GROUP {s.name} ({s.left}, {s.top}, {s.width}, {s.height})\n")
                for c in s.shapes:
                    if c.has_text_frame:
                        t = ' // '.join([p.text.strip() for p in c.text_frame.paragraphs if p.text.strip()])
                        f.write(f"    CHILD_TEXT {c.name}: {t}\n")
                    else:
                        f.write(f"    CHILD_SHAPE {c.name} ({c.shape_type})\n")
            elif s.has_text_frame:
                t = ' // '.join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])
                f.write(f"  TEXT {s.name}: {t}\n")
            else:
                f.write(f"  SHAPE {s.name} ({s.shape_type})\n")
print("Wrote slide_shape_inventory.txt")
